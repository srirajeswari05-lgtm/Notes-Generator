import streamlit as st
import fitz
from google import genai

from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from io import BytesIO


# ---------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------
st.set_page_config(
    page_title="AI College Notes Generator",
    page_icon="📚",
    layout="wide"
)


# ---------------------------------------------------------
# GEMINI CLIENT CONNECTION
# ---------------------------------------------------------
try:
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )

except KeyError:
    st.error(
        "Gemini API key was not found. "
        "Please add GEMINI_API_KEY inside "
        ".streamlit/secrets.toml."
    )
    st.stop()


# ---------------------------------------------------------
# SESSION STATE
# Stores AI notes even when Streamlit reruns the page
# ---------------------------------------------------------
if "generated_notes" not in st.session_state:
    st.session_state.generated_notes = ""


# ---------------------------------------------------------
# WORD (.docx) TEXT EXTRACTION FUNCTION
# ---------------------------------------------------------
def extract_docx_text(uploaded_file):
    document = Document(uploaded_file)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text)
# ---------------------------------------------------------
# APPLICATION HEADING
# ---------------------------------------------------------
st.title("📚 AI College Notes Generator")

st.subheader(
    "Turn your study files into smart study notes"
)

st.write(
    "Upload your college study material and generate "
    "definitions, summaries, important questions, "
    "MCQs, and quizzes using Gemini AI."
)

st.info(
    "Upload a text-based PDF and click Generate AI Study Notes."
)


# ---------------------------------------------------------
# FILE UPLOADER
# ---------------------------------------------------------
uploaded_file = st.file_uploader(
    "📂 Upload your study material",
    type=[
        "pdf",
        "ppt",
        "pptx",
        "doc",
        "docx",
        "txt",
        "xls",
        "xlsx"
    ],
    key="study_material_uploader"
)


# ---------------------------------------------------------
# FILE PROCESSING
# ---------------------------------------------------------
if uploaded_file is not None:

    st.success(
        f"{uploaded_file.name} uploaded successfully!"
    )

    file_extension = uploaded_file.name.split(".")[-1].lower()

    # -----------------------------------------------------
    # PDF PROCESSING
    # -----------------------------------------------------
    if file_extension == "pdf":

        try:
            pdf_document = fitz.open(
                stream=uploaded_file.getvalue(),
                filetype="pdf"
            )

            st.success("PDF opened successfully!")

            st.write(
                f"📄 Total Pages: {len(pdf_document)}"
            )

            text_blocks = []

            # Extract text blocks from every PDF page
            for page in pdf_document:

                blocks = page.get_text("blocks")

                for block in blocks:

                    block_text = block[4].strip()

                    if block_text:
                        text_blocks.append(block_text)

            pdf_document.close()

            all_text = "\n\n".join(text_blocks)

            # -------------------------------------------------
            # CHECK WHETHER PDF CONTAINS SELECTABLE TEXT
            # -------------------------------------------------
            if all_text.strip():

                clean_text = "\n\n".join(
                    " ".join(paragraph.split())
                    for paragraph in all_text.split("\n\n")
                    if paragraph.strip()
                )

                st.subheader("📖 Extracted Text Preview")

                st.text_area(
                    "First 1500 characters:",
                    clean_text[:1500],
                    height=300,
                    key="text_preview"
                )

                st.caption(
                    "The preview shows only the first 1500 "
                    "characters. The complete extracted text "
                    "is stored for notes generation."
                )

                # Limit very large input for the first version
                maximum_characters = 50000
                text_for_ai = clean_text[:maximum_characters]

                if len(clean_text) > maximum_characters:
                    st.warning(
                        "This PDF contains a large amount of text. "
                        "For this version, the first 50,000 "
                        "characters will be used. PDF chunking "
                        "will be added later."
                    )

                # -------------------------------------------------
                # GENERATE AI NOTES BUTTON
                # -------------------------------------------------
                if st.button(
                    "✨ Generate AI Study Notes",
                    type="primary",
                    key="generate_ai_notes_button"
                ):

                    prompt = f"""
You are an expert college study-notes creator.

Read only the study material provided below.
Do not add unrelated information.
Use clear, simple, grammatically correct English.
Make the output useful for college examinations.

Generate the response using exactly this structure:

# 📘 Chapter / Topic Title
Identify the most suitable title from the material.

## 📖 Definition
Give a clear and complete definition of the main topic.
Use one short paragraph.

## ⭐ Important Points
Provide 8 to 15 important points.
Use simple bullet points.
Do not use checkmark symbols.

## 📝 Detailed Summary
Explain the complete topic in simple English.
Use short paragraphs and preserve important concepts,
terms, formulas, or steps found in the material.

## ❓ 2-Mark Questions
Generate 5 short-answer questions.

## ❓ 5-Mark Questions
Generate 5 medium-answer questions.

## ❓ 10-Mark Questions
Generate 3 long-answer questions.

## 🧠 Multiple Choice Questions
Generate 5 MCQs.
Each MCQ must contain four options:
A, B, C, and D.

After every MCQ, write:
**Answer:** followed by the correct option and answer.

## 🎯 Quick Quiz
Generate 5 questions for student self-practice.
Do not provide answers in this section.

Important rules:
- Base the output only on the supplied material.
- Do not invent facts that are absent from the material.
- Avoid repeated questions and repeated points.
- Use professional Markdown headings.
- Keep the language easy for a college student.

STUDY MATERIAL:

{text_for_ai}
"""

                    try:
                        with st.spinner(
                            "Gemini AI is reading the material "
                            "and preparing your study notes..."
                        ):

                            response = client.models.generate_content(
                                model="gemini-3.5-flash",
                                contents=prompt
                            )

                        if response.text:

                            st.session_state.generated_notes = (
                                response.text
                            )

                            st.success(
                                "AI study notes generated successfully!"
                            )

                        else:
                            st.warning(
                                "Gemini did not return any notes. "
                                "Please try again."
                            )

                    except Exception as error:
                        st.error(
                            "AI notes generation failed."
                        )

                        st.code(str(error))

                # -------------------------------------------------
                # DISPLAY GENERATED NOTES
                # -------------------------------------------------
                if st.session_state.generated_notes:

                    st.divider()

                    st.subheader("📝 Generated Study Notes")

                    st.markdown(
                        st.session_state.generated_notes
                    )

                    # ---------------------------------------------
                    # DOWNLOAD NOTES
                    # ---------------------------------------------
                    st.download_button(
                        label="⬇️ Download Notes as Text File",
                        data=st.session_state.generated_notes,
                        file_name="AI_Generated_Study_Notes.txt",
                        mime="text/plain",
                        key="download_notes_button"
                    )

            else:
                st.warning(
                    "No selectable text was found in this PDF. "
                    "It may be a scanned or image-based PDF."
                )

        except Exception as error:
            st.error(
                "The PDF could not be processed."
            )

            st.code(str(error))

    # ---------------------------------------------------------
    # OTHER FILE TYPES
    # ---------------------------------------------------------
    else:
        st.info(
            "This file type can currently be uploaded, "
            "but text extraction is available only for PDF files. "
            "PPT, Word, TXT, and Excel extraction will be added "
            "in the next stages."
        )

else:
    st.caption(
        "Upload a PDF to begin generating AI study notes."
    )